# -*- coding: utf-8 -*-
"""
Geometry-grouped cross-validation and final explicit-equation generation for
masonry-minaret natural frequencies f1 and f2.

This public version combines:
1) the final SIMPLE linear-regression baseline, and
2) symbolic regression using PySR,
in a single reproducible script.

Both approaches use exactly the same six predictors:
    H, D0, Dtop, t, B, sqrt(E/rho)
No manually engineered ratios (H/D0, t/D0, etc.) are used in linear regression.
H is the only height-related predictor; Hc/H_silindir is not used.

Geometry groups are defined by H, D0, Dtop, t, and B. All realizations sharing
a geometry stay in the same fold. Performance is evaluated using 5-fold
GroupKFold. Final equations are re-fitted on the complete database only after
cross-validated performance evaluation.

Default repository layout
-------------------------
repo_root/
  data/minaret_frequency_database.xlsx
  results/regression/
  regression_formulas_group_cv.py

The paths can be overridden with --data and --output.
"""

import os
import argparse
from pathlib import Path
import re
import warnings
import traceback
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from sklearn.model_selection import GroupKFold
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

from pysr import PySRRegressor
from docx import Document
from docx.shared import Inches
from sympy import sympify

warnings.filterwarnings("ignore")
plt.ioff()

# ======================================================
# 1) PATHS / COMMAND-LINE OPTIONS
# ======================================================
def parse_args():
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Grouped CV for linear and symbolic minaret-frequency equations."
    )
    parser.add_argument(
        "--data",
        type=Path,
        default=script_dir / "data" / "minaret_frequency_database.xlsx",
        help="Input Excel database.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=script_dir / "results" / "regression",
        help="Directory for generated results.",
    )
    return parser.parse_args()


ARGS = parse_args()
data_path = ARGS.data.expanduser().resolve()
output_dir = ARGS.output.expanduser().resolve()

if not data_path.exists():
    raise FileNotFoundError(
        f"Input database not found: {data_path}\n"
        "Place the database at data/minaret_frequency_database.xlsx "
        "or provide a path with --data."
    )

output_dir.mkdir(parents=True, exist_ok=True)

# ======================================================
# 2) GENERAL SETTINGS
# ======================================================
RANDOM_STATE = 42
N_OUTER_SPLITS = 5

# PySR her outer fold için yeniden çalışır.
# Reviewer-odaklı group-CV yapısı korunurken arama bütçesi makul tutulmuştur.
# Eski ayar: 2000 iterations x 40 populations (çok ağır).
# Yeni ayar: 500 iterations x 12 populations. Aynı ayar outer folds ve final fit için kullanılır.
PYSR_NITERATIONS = 500
PYSR_POPULATIONS = 12
PYSR_MAXSIZE = 15
PYSR_MAXDEPTH = 6
PYSR_PARSIMONY = 1e-4

# Same geometry-group definition as the ML analysis.
GEOMETRY_COLS = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
]

TARGETS = {
    "f1 (Hz)": "f1",
    "f2 (Hz)": "f2",
}

# ======================================================
# 3) DATA-CLEANING / GROUP HELPERS
# ======================================================
def clean_numeric_series(series):
    return pd.to_numeric(
        series.astype(str)
        .str.replace(",", ".", regex=False)
        .str.replace("[", "", regex=False)
        .str.replace("]", "", regex=False)
        .str.strip(),
        errors="coerce",
    )


def clean_name(text):
    return (
        str(text)
        .replace(" ", "_")
        .replace("(", "")
        .replace(")", "")
        .replace("/", "_")
    )


def make_geometry_groups(dataframe, geometry_cols, decimals=8):
    """ML koduyla uyumlu, satır sırasından bağımsız G1...Gn geometry IDs."""
    missing = [c for c in geometry_cols if c not in dataframe.columns]
    if missing:
        raise ValueError(
            "Geometry group oluşturmak için gerekli sütunlar eksik: "
            + ", ".join(missing)
        )

    geom = dataframe[geometry_cols].apply(pd.to_numeric, errors="coerce").round(decimals)
    if geom.isna().any().any():
        raise ValueError("Geometry columns içinde NaN var; group oluşturulamıyor.")

    tuples = [tuple(row) for row in geom.to_numpy()]
    unique_sorted = sorted(set(tuples))
    tuple_to_group = {key: f"G{i + 1}" for i, key in enumerate(unique_sorted)}

    groups = pd.Series(
        [tuple_to_group[key] for key in tuples],
        index=dataframe.index,
        name="Geometry_Group",
    )
    keys = pd.Series(
        ["_".join(f"{v:g}" for v in key) for key in tuples],
        index=dataframe.index,
        name="Geometry_Key",
    )
    return groups, keys


# ======================================================
# 4) METRICS
# ======================================================
def smape_percent(y_true, y_pred):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    denom = np.abs(y_true) + np.abs(y_pred)
    valid = denom > 0
    if not np.any(valid):
        return 0.0
    return float(
        np.mean(2.0 * np.abs(y_pred[valid] - y_true[valid]) / denom[valid]) * 100.0
    )


def within_relative_band_percent(y_true, y_pred, band=0.20):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    valid = np.abs(y_true) > 1e-12
    if not np.any(valid):
        return np.nan
    rel_err = np.abs(y_pred[valid] - y_true[valid]) / np.abs(y_true[valid])
    return float(np.mean(rel_err <= band) * 100.0)


def calc_metrics(y_true, y_pred):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)

    mae = mean_absolute_error(y_true, y_pred)
    rmse = np.sqrt(mean_squared_error(y_true, y_pred))
    r2 = r2_score(y_true, y_pred)
    smape = smape_percent(y_true, y_pred)
    within20 = within_relative_band_percent(y_true, y_pred, band=0.20)

    if np.std(y_true) == 0 or np.std(y_pred) == 0:
        r = np.nan
    else:
        r = np.corrcoef(y_true, y_pred)[0, 1]

    return {
        "R2": float(r2),
        "MAE": float(mae),
        "RMSE": float(rmse),
        "SMAPE_pct": float(smape),
        "CC": float(r) if pd.notna(r) else np.nan,
        "Within_20pct_pct": float(within20) if pd.notna(within20) else np.nan,
    }


def mean_sd(series):
    s = pd.Series(series, dtype=float)
    return float(s.mean()), float(s.std(ddof=1))


# ======================================================
# 5) FORMULA HELPERS
# ======================================================
def linear_formula(coefs, intercept, symbols, yname, tol=1e-10):
    terms = [f"{intercept:.8g}"]
    for c, s in zip(coefs, symbols):
        if abs(c) < tol:
            continue
        sign = " + " if c >= 0 else " - "
        terms.append(f"{sign}{abs(c):.8g}*{s}")
    return f"{yname} = " + "".join(terms)


def replace_x_with_symbols(expr_text, x_symbol_map):
    # x10 içinde x1'in yanlış replace edilmesini önlemek için uzun isimlerden başla.
    out = str(expr_text)
    for old in sorted(x_symbol_map.keys(), key=len, reverse=True):
        out = re.sub(
            rf"(?<![A-Za-z0-9_]){re.escape(old)}(?![A-Za-z0-9_])",
            x_symbol_map[old],
            out,
        )
    return out


def simplify_formula_text(expr_text, x_symbol_map):
    expr_text = replace_x_with_symbols(expr_text, x_symbol_map)
    try:
        expr = sympify(expr_text)
        expr = expr.simplify()
        return str(expr)
    except Exception:
        return expr_text


def count_used_variables(formula_text, allowed_symbols):
    used = []
    for s in allowed_symbols:
        pattern = rf"(?<![A-Za-z0-9_]){re.escape(s)}(?![A-Za-z0-9_])"
        if re.search(pattern, formula_text):
            used.append(s)
    return sorted(set(used))


def build_pysr(random_state):
    return PySRRegressor(
        niterations=PYSR_NITERATIONS,
        populations=PYSR_POPULATIONS,
        maxsize=PYSR_MAXSIZE,
        maxdepth=PYSR_MAXDEPTH,
        binary_operators=["+", "-", "*", "/"],
        unary_operators=["sqrt", "log"],
        constraints={
            "sqrt": 1,
            "log": 1,
            "/": (-1, 9),
        },
        nested_constraints={
            "sqrt": {"sqrt": 0},
            "log": {"log": 0},
        },
        parsimony=PYSR_PARSIMONY,
        model_selection="best",
        progress=True,
        turbo=True,
        random_state=random_state,
    )


def get_selected_pysr_formula(model_pysr, x_symbol_map, allowed_symbols):
    """
    Modelin predict() için seçtiği ifadeyi olabildiğince doğrudan alır.
    PySR sürümleri arasındaki API farklarına karşı fallback içerir.
    """
    raw_expr = None
    complexity = np.nan
    loss = np.nan

    try:
        raw_expr = str(model_pysr.sympy())
    except Exception:
        pass

    if raw_expr is None:
        try:
            best_row = model_pysr.get_best()
            raw_expr = str(best_row["sympy_format"])
            complexity = best_row.get("complexity", np.nan)
            loss = best_row.get("loss", np.nan)
        except Exception:
            pass

    if raw_expr is None:
        eq_df = model_pysr.equations_.copy()
        if "score" in eq_df.columns:
            best_idx = eq_df["score"].idxmax()
        else:
            best_idx = eq_df["loss"].idxmin()
        best_row = eq_df.loc[best_idx]
        raw_expr = str(best_row["sympy_format"])
        complexity = best_row.get("complexity", np.nan)
        loss = best_row.get("loss", np.nan)

    formula_text = simplify_formula_text(raw_expr, x_symbol_map)
    used_vars = count_used_variables(formula_text, allowed_symbols)

    return formula_text, complexity, loss, used_vars


# ======================================================
# 6) PARITY PLOTS / DIAGNOSTICS
# ======================================================
def save_parity_plot(
    y_true,
    y_pred,
    fold_ids,
    model_name,
    target_name,
    metrics,
    save_path,
):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    fold_ids = np.asarray(fold_ids, dtype=int)

    fig, ax = plt.subplots(figsize=(7.0, 7.0))

    unique_folds = sorted(np.unique(fold_ids))
    cmap = plt.get_cmap("tab10")
    for i, fold in enumerate(unique_folds):
        mask = fold_ids == fold
        ax.scatter(
            y_true[mask],
            y_pred[mask],
            alpha=0.72,
            s=34,
            label=f"Outer fold {fold}",
            color=cmap(i % 10),
        )

    min_val = min(np.min(y_true), np.min(y_pred))
    max_val = max(np.max(y_true), np.max(y_pred))
    lower = max(0.0, min_val * 0.95)
    upper = max_val * 1.05
    line_x = np.linspace(lower, upper, 250)

    # 1:1 ideal line
    ax.plot(line_x, line_x, linestyle="--", linewidth=1.6, label="1:1 line")

    # Reviewer-oriented ±20% error bands
    ax.plot(line_x, 1.20 * line_x, linestyle=":", linewidth=1.3, label="±20% bounds")
    ax.plot(line_x, 0.80 * line_x, linestyle=":", linewidth=1.3)

    text_box = (
        f"Pooled OOF R² = {metrics['R2']:.4f}\n"
        f"MAE = {metrics['MAE']:.4f} Hz\n"
        f"RMSE = {metrics['RMSE']:.4f} Hz\n"
        f"SMAPE = {metrics['SMAPE_pct']:.2f}%\n"
        f"Within ±20% = {metrics['Within_20pct_pct']:.1f}%"
    )

    ax.text(
        0.03,
        0.97,
        text_box,
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=13.5,
        bbox=dict(facecolor="white", alpha=0.88),
    )

    ax.set_xlabel(f"FE reference {target_name}", fontsize=15)
    ax.set_ylabel(f"Grouped outer-test prediction {target_name}", fontsize=15)
    ax.set_title(f"{target_name} – {model_name}", fontsize=16)
    ax.set_xlim(lower, upper)
    ax.set_ylim(lower, upper)
    ax.set_aspect("equal", adjustable="box")
    ax.tick_params(labelsize=13)
    ax.legend(fontsize=11, loc="lower right")

    fig.tight_layout()
    fig.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def create_combined_figure(image_paths, labels, output_path, main_title=""):
    n = len(image_paths)
    if n == 0:
        return None

    cols = 2
    rows = int(np.ceil(n / cols))
    fig, axes = plt.subplots(rows, cols, figsize=(13, 6.2 * rows))
    axes_flat = np.atleast_1d(axes).flatten()

    for i, ax in enumerate(axes_flat):
        if i < n:
            img = plt.imread(image_paths[i])
            ax.imshow(img)
            ax.axis("off")
            ax.text(
                0.5,
                -0.045,
                labels[i],
                transform=ax.transAxes,
                ha="center",
                va="top",
                fontsize=14,
            )
        else:
            ax.axis("off")

    if main_title:
        fig.suptitle(main_title, fontsize=19)

    plt.tight_layout(rect=[0, 0.02, 1, 0.96])
    plt.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return output_path


def high_frequency_tail_metrics(y_true, y_pred, model_name, target_name, quantile=0.90):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)

    threshold = float(np.quantile(y_true, quantile))
    mask = y_true >= threshold
    if not np.any(mask):
        return None

    yt = y_true[mask]
    yp = y_pred[mask]
    signed_error = yt - yp  # positive => underprediction

    return {
        "Target": target_name,
        "Model": model_name,
        "Quantile_Threshold": quantile,
        "Frequency_Threshold_Hz": threshold,
        "n_Cases": int(mask.sum()),
        "Mean_Signed_Error_ActualMinusPred_Hz": float(np.mean(signed_error)),
        "MAE_Hz": float(mean_absolute_error(yt, yp)),
        "RMSE_Hz": float(np.sqrt(mean_squared_error(yt, yp))),
        "SMAPE_pct": float(smape_percent(yt, yp)),
        "Within_20pct_pct": float(within_relative_band_percent(yt, yp, 0.20)),
    }


# ======================================================
# 7) LOAD / CLEAN DATA / DERIVED MATERIAL PARAMETER
# ======================================================
df = pd.read_excel(data_path)
df.columns = [str(c).strip() for c in df.columns]

raw_needed_cols = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
    "E (MPa)",
    "Ro (kg/m3)",
    "f1 (Hz)",
    "f2 (Hz)",
]

missing = [c for c in raw_needed_cols if c not in df.columns]
if missing:
    raise ValueError("Eksik gerekli sütunlar: " + ", ".join(missing))

for col in raw_needed_cols:
    df[col] = clean_numeric_series(df[col])

# Keep only complete rows required by the published analysis.
df = df.dropna(subset=raw_needed_cols).reset_index(drop=True)

if (df["Ro (kg/m3)"] <= 0).any() or (df["E (MPa)"] < 0).any():
    raise ValueError("sqrt(E/rho) requires E >= 0 and rho > 0.")

df["sqrt_E_div_rho"] = np.sqrt(df["E (MPa)"] / df["Ro (kg/m3)"])

feature_cols_symbolic = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
    "sqrt_E_div_rho",
]

feature_cols_linear = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
    "sqrt_E_div_rho",
]

symbol_map_symbolic = {
    "H_m": "H",
    "D0_m": "D0",
    "Dtop_m": "Dtop",
    "t_m": "t",
    "B_kaide_m": "B",
    "sqrt_E_div_rho": "sqrtE_rho",
}

symbol_map_linear = {
    "H_m": "H",
    "D0_m": "D0",
    "Dtop_m": "Dtop",
    "t_m": "t",
    "B_kaide_m": "B",
    "sqrt_E_div_rho": "sqrtE_rho",
}

feature_symbols_linear = [symbol_map_linear[c] for c in feature_cols_linear]
feature_symbols_symbolic = [symbol_map_symbolic[c] for c in feature_cols_symbolic]

x_symbol_map = {
    "x0": "H",
    "x1": "D0",
    "x2": "Dtop",
    "x3": "t",
    "x4": "B",
    "x5": "sqrtE_rho",
}

# Geometry groups: identical to the ML analysis.
groups_full, geometry_keys = make_geometry_groups(df, GEOMETRY_COLS)
df["Geometry_Group"] = groups_full
df["Geometry_Key"] = geometry_keys

n_groups = df["Geometry_Group"].nunique()
if n_groups < N_OUTER_SPLITS:
    raise ValueError(
        f"{N_OUTER_SPLITS}-fold GroupKFold için en az {N_OUTER_SPLITS} geometry group gerekli; "
        f"bulunan: {n_groups}"
    )

print("Temizlenmiş veri boyutu:", df.shape)
print("Unique geometry group sayısı:", n_groups)
print("Geometry group definition:", ", ".join(GEOMETRY_COLS))

# ======================================================
# 8) GROUPED OUTER CV: LINEAR REGRESSION
# ======================================================
def evaluate_linear_group_cv(X_df, y, groups, target_name):
    outer_cv = GroupKFold(n_splits=N_OUTER_SPLITS)

    oof_pred = np.full(len(y), np.nan, dtype=float)
    fold_ids = np.full(len(y), np.nan, dtype=float)
    fold_rows = []
    formula_rows = []

    for fold_no, (train_idx, test_idx) in enumerate(
        outer_cv.split(X_df, y, groups), start=1
    ):
        X_train = X_df.iloc[train_idx].copy()
        X_test = X_df.iloc[test_idx].copy()
        y_train = y.iloc[train_idx].copy()
        y_test = y.iloc[test_idx].copy()
        g_train = groups.iloc[train_idx].copy()
        g_test = groups.iloc[test_idx].copy()

        overlap = set(g_train.astype(str)).intersection(set(g_test.astype(str)))
        if overlap:
            raise RuntimeError(
                f"Linear Regression | {target_name} | fold {fold_no}: group leakage: {sorted(overlap)}"
            )

        model = LinearRegression()
        model.fit(X_train, y_train)

        train_pred = model.predict(X_train)
        test_pred = model.predict(X_test)

        oof_pred[test_idx] = test_pred
        fold_ids[test_idx] = fold_no

        train_m = calc_metrics(y_train, train_pred)
        test_m = calc_metrics(y_test, test_pred)

        fold_formula = linear_formula(
            model.coef_, model.intercept_, feature_symbols_linear, target_name
        )

        fold_rows.append({
            "Target": target_name,
            "Model": "Linear Regression",
            "Outer_Fold": fold_no,
            "Train_n": len(train_idx),
            "Test_n": len(test_idx),
            "Train_Groups": g_train.nunique(),
            "Test_Groups": g_test.nunique(),
            "Group_Overlap_n": len(overlap),
            "Train_R2": train_m["R2"],
            "Train_MAE_Hz": train_m["MAE"],
            "Train_RMSE_Hz": train_m["RMSE"],
            "Train_SMAPE_pct": train_m["SMAPE_pct"],
            "Train_Within_20pct_pct": train_m["Within_20pct_pct"],
            "Test_R2": test_m["R2"],
            "Test_MAE_Hz": test_m["MAE"],
            "Test_RMSE_Hz": test_m["RMSE"],
            "Test_SMAPE_pct": test_m["SMAPE_pct"],
            "Test_Within_20pct_pct": test_m["Within_20pct_pct"],
        })

        formula_rows.append({
            "Target": target_name,
            "Model": "Linear Regression",
            "Outer_Fold": fold_no,
            "Formula": fold_formula,
        })

        print(
            f"{target_name} | Linear Regression | Outer fold {fold_no}/{N_OUTER_SPLITS} | "
            f"train groups={g_train.nunique()} | test groups={g_test.nunique()} | "
            f"R²={test_m['R2']:.4f} | SMAPE={test_m['SMAPE_pct']:.2f}%"
        )

    if np.isnan(oof_pred).any() or np.isnan(fold_ids).any():
        raise RuntimeError(f"{target_name} Linear Regression: OOF prediction eksik.")

    return (
        oof_pred,
        fold_ids.astype(int),
        pd.DataFrame(fold_rows),
        pd.DataFrame(formula_rows),
    )


# ======================================================
# 9) GROUPED OUTER CV: SYMBOLIC REGRESSION (PySR)
# ======================================================
def evaluate_pysr_group_cv(X_df, y, groups, target_name):
    outer_cv = GroupKFold(n_splits=N_OUTER_SPLITS)

    oof_pred = np.full(len(y), np.nan, dtype=float)
    fold_ids = np.full(len(y), np.nan, dtype=float)
    fold_rows = []
    formula_rows = []

    for fold_no, (train_idx, test_idx) in enumerate(
        outer_cv.split(X_df, y, groups), start=1
    ):
        X_train = X_df.iloc[train_idx].copy()
        X_test = X_df.iloc[test_idx].copy()
        y_train = y.iloc[train_idx].copy()
        y_test = y.iloc[test_idx].copy()
        g_train = groups.iloc[train_idx].copy()
        g_test = groups.iloc[test_idx].copy()

        overlap = set(g_train.astype(str)).intersection(set(g_test.astype(str)))
        if overlap:
            raise RuntimeError(
                f"PySR | {target_name} | fold {fold_no}: group leakage: {sorted(overlap)}"
            )

        model = build_pysr(RANDOM_STATE + fold_no)
        print(
            f"{target_name} | PySR | Outer fold {fold_no}/{N_OUTER_SPLITS} fit başlıyor..."
        )
        model.fit(X_train.to_numpy(), y_train.to_numpy())

        train_pred = model.predict(X_train.to_numpy())
        test_pred = model.predict(X_test.to_numpy())

        oof_pred[test_idx] = test_pred
        fold_ids[test_idx] = fold_no

        train_m = calc_metrics(y_train, train_pred)
        test_m = calc_metrics(y_test, test_pred)

        try:
            formula_text, complexity, loss, used_vars = get_selected_pysr_formula(
                model,
                x_symbol_map,
                feature_symbols_symbolic,
            )
            formula_text = f"{target_name} = {formula_text}"
        except Exception as exc:
            formula_text = f"Formula extraction failed: {exc}"
            complexity = np.nan
            loss = np.nan
            used_vars = []

        fold_rows.append({
            "Target": target_name,
            "Model": "Symbolic Regression (PySR)",
            "Outer_Fold": fold_no,
            "Train_n": len(train_idx),
            "Test_n": len(test_idx),
            "Train_Groups": g_train.nunique(),
            "Test_Groups": g_test.nunique(),
            "Group_Overlap_n": len(overlap),
            "Train_R2": train_m["R2"],
            "Train_MAE_Hz": train_m["MAE"],
            "Train_RMSE_Hz": train_m["RMSE"],
            "Train_SMAPE_pct": train_m["SMAPE_pct"],
            "Train_Within_20pct_pct": train_m["Within_20pct_pct"],
            "Test_R2": test_m["R2"],
            "Test_MAE_Hz": test_m["MAE"],
            "Test_RMSE_Hz": test_m["RMSE"],
            "Test_SMAPE_pct": test_m["SMAPE_pct"],
            "Test_Within_20pct_pct": test_m["Within_20pct_pct"],
            "Selected_Complexity": complexity,
            "Selected_Loss": loss,
            "Used_Variable_Count": len(used_vars),
            "Used_Variables": ", ".join(used_vars),
        })

        formula_rows.append({
            "Target": target_name,
            "Model": "Symbolic Regression (PySR)",
            "Outer_Fold": fold_no,
            "Formula": formula_text,
            "Complexity": complexity,
            "Loss": loss,
            "Used_Variables": ", ".join(used_vars),
        })

        print(
            f"{target_name} | Symbolic Regression | Outer fold {fold_no}/{N_OUTER_SPLITS} | "
            f"train groups={g_train.nunique()} | test groups={g_test.nunique()} | "
            f"R²={test_m['R2']:.4f} | SMAPE={test_m['SMAPE_pct']:.2f}%"
        )

    if np.isnan(oof_pred).any() or np.isnan(fold_ids).any():
        raise RuntimeError(f"{target_name} PySR: OOF prediction eksik.")

    return (
        oof_pred,
        fold_ids.astype(int),
        pd.DataFrame(fold_rows),
        pd.DataFrame(formula_rows),
    )


# ======================================================
# 10) PERFORMANCE SUMMARY
# ======================================================
def summarize_fold_performance(fold_df, pooled_metrics):
    def ms(col):
        return mean_sd(fold_df[col])

    tr_r2_m, tr_r2_sd = ms("Train_R2")
    tr_mae_m, tr_mae_sd = ms("Train_MAE_Hz")
    tr_rmse_m, tr_rmse_sd = ms("Train_RMSE_Hz")
    tr_smape_m, tr_smape_sd = ms("Train_SMAPE_pct")

    te_r2_m, te_r2_sd = ms("Test_R2")
    te_mae_m, te_mae_sd = ms("Test_MAE_Hz")
    te_rmse_m, te_rmse_sd = ms("Test_RMSE_Hz")
    te_smape_m, te_smape_sd = ms("Test_SMAPE_pct")
    te_20_m, te_20_sd = ms("Test_Within_20pct_pct")

    return {
        "Train_R2_Mean": tr_r2_m,
        "Train_R2_SD": tr_r2_sd,
        "Train_MAE_Mean_Hz": tr_mae_m,
        "Train_MAE_SD_Hz": tr_mae_sd,
        "Train_RMSE_Mean_Hz": tr_rmse_m,
        "Train_RMSE_SD_Hz": tr_rmse_sd,
        "Train_SMAPE_Mean_pct": tr_smape_m,
        "Train_SMAPE_SD_pct": tr_smape_sd,
        "Test_R2_Mean": te_r2_m,
        "Test_R2_SD": te_r2_sd,
        "Test_MAE_Mean_Hz": te_mae_m,
        "Test_MAE_SD_Hz": te_mae_sd,
        "Test_RMSE_Mean_Hz": te_rmse_m,
        "Test_RMSE_SD_Hz": te_rmse_sd,
        "Test_SMAPE_Mean_pct": te_smape_m,
        "Test_SMAPE_SD_pct": te_smape_sd,
        "Test_Within_20pct_Mean_pct": te_20_m,
        "Test_Within_20pct_SD_pct": te_20_sd,
        "Pooled_OOF_R2": pooled_metrics["R2"],
        "Pooled_OOF_MAE_Hz": pooled_metrics["MAE"],
        "Pooled_OOF_RMSE_Hz": pooled_metrics["RMSE"],
        "Pooled_OOF_SMAPE_pct": pooled_metrics["SMAPE_pct"],
        "Pooled_OOF_CC": pooled_metrics["CC"],
        "Pooled_OOF_Within_20pct_pct": pooled_metrics["Within_20pct_pct"],
    }


def make_paper_ready_table(summary_df):
    rows = []
    for _, r in summary_df.iterrows():
        rows.append({
            "Target": r["Target"],
            "Model": r["Model"],
            "Training R2 (mean ± SD)": f"{r['Train_R2_Mean']:.4f} ± {r['Train_R2_SD']:.4f}",
            "Test R2 (mean ± SD)": f"{r['Test_R2_Mean']:.4f} ± {r['Test_R2_SD']:.4f}",
            "Training MAE Hz (mean ± SD)": f"{r['Train_MAE_Mean_Hz']:.4f} ± {r['Train_MAE_SD_Hz']:.4f}",
            "Test MAE Hz (mean ± SD)": f"{r['Test_MAE_Mean_Hz']:.4f} ± {r['Test_MAE_SD_Hz']:.4f}",
            "Training RMSE Hz (mean ± SD)": f"{r['Train_RMSE_Mean_Hz']:.4f} ± {r['Train_RMSE_SD_Hz']:.4f}",
            "Test RMSE Hz (mean ± SD)": f"{r['Test_RMSE_Mean_Hz']:.4f} ± {r['Test_RMSE_SD_Hz']:.4f}",
            "Test SMAPE % (mean ± SD)": f"{r['Test_SMAPE_Mean_pct']:.2f} ± {r['Test_SMAPE_SD_pct']:.2f}",
            "Pooled OOF R2": f"{r['Pooled_OOF_R2']:.4f}",
            "Pooled within ±20% (%)": f"{r['Pooled_OOF_Within_20pct_pct']:.1f}",
        })
    return pd.DataFrame(rows)


# ======================================================
# 11) WORD REPORT
# ======================================================
doc = Document()
doc.add_heading("Group-Based Cross-Validation ile Frekans Formülleri", level=1)
doc.add_paragraph(
    "Simple linear and symbolic regression performance was evaluated using geometry-grouped "
    "five-fold cross-validation. All realizations sharing the same geometry were kept "
    "within the same fold, so the reported test performance represents generalization "
    "to previously unseen geometry groups within the adopted numerical database."
)
doc.add_paragraph("Geometry group definition: " + ", ".join(GEOMETRY_COLS) + ".")
doc.add_paragraph(
    f"The cleaned dataset contains {len(df)} realizations and {n_groups} unique geometry groups."
)
doc.add_paragraph(
    "Parity plots are based on pooled held-out outer-test predictions and include the "
    "ideal 1:1 line and ±20% relative-error bounds. Final equations are re-fitted using "
    "the complete dataset only after cross-validated performance evaluation."
)

# ======================================================
# 12) MAIN ANALYSIS LOOP
# ======================================================
all_summary_rows = []
all_fold_rows = []
all_oof_rows = []
all_outer_formula_rows = []
all_final_formula_rows = []
all_final_pysr_equations = []
all_high_freq_rows = []
all_plot_paths = {}

for target_col, label in TARGETS.items():
    print("\n" + "=" * 90)
    print(f"{target_col} için GROUP-BASED formula evaluation başlıyor...")
    print("=" * 90)

    target_df = df.dropna(subset=[target_col]).reset_index(drop=True)
    y = target_df[target_col].astype(float)
    groups = target_df["Geometry_Group"].astype(str)

    X_linear = target_df[feature_cols_linear].copy()
    X_symbolic = target_df[feature_cols_symbolic].copy()

    if groups.nunique() < N_OUTER_SPLITS:
        raise ValueError(
            f"{target_col}: yalnızca {groups.nunique()} unique geometry group var; "
            f"{N_OUTER_SPLITS} fold uygulanamaz."
        )

    # --------------------------------------------------
    # A) LINEAR REGRESSION - GROUP CV
    # --------------------------------------------------
    (
        pred_lin,
        fold_lin,
        fold_df_lin,
        formula_df_lin,
    ) = evaluate_linear_group_cv(
        X_linear,
        y,
        groups,
        target_col,
    )

    pooled_lin = calc_metrics(y, pred_lin)
    summary_lin = summarize_fold_performance(fold_df_lin, pooled_lin)
    summary_lin.update({"Target": target_col, "Model": "Linear Regression"})
    all_summary_rows.append(summary_lin)
    all_fold_rows.append(fold_df_lin)
    all_outer_formula_rows.append(formula_df_lin)

    all_oof_rows.append(pd.DataFrame({
        "Target": target_col,
        "Model": "Linear Regression",
        "Row_Index": target_df.index,
        "Geometry_Group": groups.values,
        "Outer_Fold": fold_lin,
        "Actual_Hz": y.values,
        "Grouped_OuterTest_Predicted_Hz": pred_lin,
        "Residual_ActualMinusPred_Hz": y.values - pred_lin,
    }))

    hf_lin = high_frequency_tail_metrics(
        y, pred_lin, "Linear Regression", target_col, quantile=0.90
    )
    if hf_lin is not None:
        all_high_freq_rows.append(hf_lin)

    lin_plot = os.path.join(
        output_dir,
        f"Parity_GroupCV_{clean_name(target_col)}_LinearRegression.png",
    )
    save_parity_plot(
        y,
        pred_lin,
        fold_lin,
        "Linear Regression",
        target_col,
        pooled_lin,
        lin_plot,
    )

    # --------------------------------------------------
    # B) SYMBOLIC REGRESSION - GROUP CV
    # --------------------------------------------------
    try:
        (
            pred_pysr,
            fold_pysr,
            fold_df_pysr,
            formula_df_pysr,
        ) = evaluate_pysr_group_cv(
            X_symbolic,
            y,
            groups,
            target_col,
        )
    except Exception:
        print(f"PySR group CV başarısız: {target_col}")
        traceback.print_exc()
        raise

    pooled_pysr = calc_metrics(y, pred_pysr)
    summary_pysr = summarize_fold_performance(fold_df_pysr, pooled_pysr)
    summary_pysr.update({"Target": target_col, "Model": "Symbolic Regression (PySR)"})
    all_summary_rows.append(summary_pysr)
    all_fold_rows.append(fold_df_pysr)
    all_outer_formula_rows.append(formula_df_pysr)

    all_oof_rows.append(pd.DataFrame({
        "Target": target_col,
        "Model": "Symbolic Regression (PySR)",
        "Row_Index": target_df.index,
        "Geometry_Group": groups.values,
        "Outer_Fold": fold_pysr,
        "Actual_Hz": y.values,
        "Grouped_OuterTest_Predicted_Hz": pred_pysr,
        "Residual_ActualMinusPred_Hz": y.values - pred_pysr,
    }))

    hf_pysr = high_frequency_tail_metrics(
        y, pred_pysr, "Symbolic Regression (PySR)", target_col, quantile=0.90
    )
    if hf_pysr is not None:
        all_high_freq_rows.append(hf_pysr)

    pysr_plot = os.path.join(
        output_dir,
        f"Parity_GroupCV_{clean_name(target_col)}_PySR.png",
    )
    save_parity_plot(
        y,
        pred_pysr,
        fold_pysr,
        "Symbolic Regression (PySR)",
        target_col,
        pooled_pysr,
        pysr_plot,
    )

    combined_plot = os.path.join(
        output_dir,
        f"Combined_GroupCV_Parity_{clean_name(target_col)}.png",
    )
    create_combined_figure(
        [lin_plot, pysr_plot],
        ["a) Linear Regression", "b) Symbolic Regression (PySR)"],
        combined_plot,
        main_title=f"{target_col} – grouped outer-test parity plots",
    )
    all_plot_paths[target_col] = combined_plot

    # --------------------------------------------------
    # C) FINAL FULL-DATA LINEAR FORMULA
    # --------------------------------------------------
    final_lin = LinearRegression()
    final_lin.fit(X_linear, y)
    final_lin_formula = linear_formula(
        final_lin.coef_, final_lin.intercept_, feature_symbols_linear, label
    )
    final_lin_train_pred = final_lin.predict(X_linear)
    final_lin_fit_metrics = calc_metrics(y, final_lin_train_pred)

    all_final_formula_rows.append({
        "Target": target_col,
        "Model": "Linear Regression",
        "Final_Formula": final_lin_formula,
        "FullData_Fit_R2_NOT_CV": final_lin_fit_metrics["R2"],
        "FullData_Fit_MAE_Hz_NOT_CV": final_lin_fit_metrics["MAE"],
        "FullData_Fit_RMSE_Hz_NOT_CV": final_lin_fit_metrics["RMSE"],
        "Note": "Use group-CV metrics for performance claims; full-data metrics are fit diagnostics only.",
    })

    # --------------------------------------------------
    # D) FINAL FULL-DATA PySR FORMULA
    # --------------------------------------------------
    print(f"{target_col} | FINAL PySR full-data fit başlıyor...")
    final_pysr = build_pysr(RANDOM_STATE)
    final_pysr.fit(X_symbolic.to_numpy(), y.to_numpy())

    final_pysr_pred = final_pysr.predict(X_symbolic.to_numpy())
    final_pysr_fit_metrics = calc_metrics(y, final_pysr_pred)

    final_formula_only, final_complexity, final_loss, final_used_vars = get_selected_pysr_formula(
        final_pysr,
        x_symbol_map,
        feature_symbols_symbolic,
    )
    final_pysr_formula = f"{label} = {final_formula_only}"

    all_final_formula_rows.append({
        "Target": target_col,
        "Model": "Symbolic Regression (PySR)",
        "Final_Formula": final_pysr_formula,
        "Complexity": final_complexity,
        "Loss": final_loss,
        "Used_Variable_Count": len(final_used_vars),
        "Used_Variables": ", ".join(final_used_vars),
        "FullData_Fit_R2_NOT_CV": final_pysr_fit_metrics["R2"],
        "FullData_Fit_MAE_Hz_NOT_CV": final_pysr_fit_metrics["MAE"],
        "FullData_Fit_RMSE_Hz_NOT_CV": final_pysr_fit_metrics["RMSE"],
        "Note": "Use group-CV metrics for performance claims; full-data metrics are fit diagnostics only.",
    })

    # Tüm final PySR candidate equations'i reproducibility için kaydet.
    try:
        eq_df = final_pysr.equations_.copy()
        eq_df = eq_df.reset_index().rename(columns={"index": "Equation_Index"})
        if "sympy_format" in eq_df.columns:
            eq_df["Formula_Readable"] = [
                simplify_formula_text(str(v), x_symbol_map)
                for v in eq_df["sympy_format"]
            ]
            eq_df["sympy_format"] = eq_df["sympy_format"].astype(str)
        eq_df.insert(0, "Target", target_col)
        all_final_pysr_equations.append(eq_df)
    except Exception as exc:
        print(f"Final PySR candidate equations export edilemedi: {target_col} | {exc}")

    # --------------------------------------------------
    # E) WORD TARGET SECTION
    # --------------------------------------------------
    doc.add_heading(f"{label} Results", level=2)

    target_summary = pd.DataFrame([summary_lin, summary_pysr])
    target_paper = make_paper_ready_table(target_summary)

    table = doc.add_table(rows=1, cols=8)
    hdr = table.rows[0].cells
    headers = [
        "Model",
        "Training R²",
        "Test R²",
        "Training MAE",
        "Test MAE",
        "Training RMSE",
        "Test RMSE",
        "Test SMAPE",
    ]
    for i, h in enumerate(headers):
        hdr[i].text = h

    for _, row in target_summary.iterrows():
        c = table.add_row().cells
        c[0].text = row["Model"]
        c[1].text = f"{row['Train_R2_Mean']:.4f} ± {row['Train_R2_SD']:.4f}"
        c[2].text = f"{row['Test_R2_Mean']:.4f} ± {row['Test_R2_SD']:.4f}"
        c[3].text = f"{row['Train_MAE_Mean_Hz']:.4f} ± {row['Train_MAE_SD_Hz']:.4f}"
        c[4].text = f"{row['Test_MAE_Mean_Hz']:.4f} ± {row['Test_MAE_SD_Hz']:.4f}"
        c[5].text = f"{row['Train_RMSE_Mean_Hz']:.4f} ± {row['Train_RMSE_SD_Hz']:.4f}"
        c[6].text = f"{row['Test_RMSE_Mean_Hz']:.4f} ± {row['Test_RMSE_SD_Hz']:.4f}"
        c[7].text = f"{row['Test_SMAPE_Mean_pct']:.2f} ± {row['Test_SMAPE_SD_pct']:.2f}%"

    doc.add_heading("Final full-data equations", level=3)
    doc.add_paragraph(
        "The following equations were re-fitted using the complete dataset only after "
        "group-based cross-validation. Their predictive performance should be reported "
        "using the grouped outer-test results above, not the full-data fit statistics."
    )
    doc.add_paragraph(final_lin_formula)
    doc.add_paragraph(final_pysr_formula)

    doc.add_heading("Grouped outer-test parity plots", level=3)
    doc.add_picture(combined_plot, width=Inches(6.5))

    print(f"\n{label} FINAL linear formula:")
    print(final_lin_formula)
    print(f"\n{label} FINAL symbolic formula:")
    print(final_pysr_formula)

# ======================================================
# 13) COMBINE OUTPUTS
# ======================================================
summary_df = pd.DataFrame(all_summary_rows)
summary_df = summary_df.sort_values(["Target", "Test_R2_Mean"], ascending=[True, False]).reset_index(drop=True)

paper_ready_df = make_paper_ready_table(summary_df)
fold_results_df = pd.concat(all_fold_rows, ignore_index=True)
oof_predictions_df = pd.concat(all_oof_rows, ignore_index=True)
outer_formulas_df = pd.concat(all_outer_formula_rows, ignore_index=True)
final_formulas_df = pd.DataFrame(all_final_formula_rows)
high_freq_df = pd.DataFrame(all_high_freq_rows)

if all_final_pysr_equations:
    final_pysr_equations_df = pd.concat(all_final_pysr_equations, ignore_index=True)
else:
    final_pysr_equations_df = pd.DataFrame()

settings_df = pd.DataFrame({
    "Setting": [
        "Random_State",
        "Outer_GroupKFold_Splits",
        "Geometry_Group_Columns",
        "Unique_Geometry_Groups",
        "Rows_After_Cleaning",
        "PySR_niterations",
        "PySR_populations",
        "PySR_maxsize",
        "PySR_maxdepth",
        "PySR_parsimony",
        "Symbolic_Features",
        "Linear_Features",
    ],
    "Value": [
        RANDOM_STATE,
        N_OUTER_SPLITS,
        ", ".join(GEOMETRY_COLS),
        n_groups,
        len(df),
        PYSR_NITERATIONS,
        PYSR_POPULATIONS,
        PYSR_MAXSIZE,
        PYSR_MAXDEPTH,
        PYSR_PARSIMONY,
        ", ".join(feature_cols_symbolic),
        ", ".join(feature_cols_linear),
    ],
})

# ======================================================
# 14) SAVE EXCEL OUTPUT
# ======================================================
excel_path = os.path.join(output_dir, "Frekans_Formul_GroupCV_Sonuclari.xlsx")

with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
    # Makaleye doğrudan aktarılabilecek ana tablo
    paper_ready_df.to_excel(writer, sheet_name="Paper_Ready_Performance", index=False)

    # Tüm sayısal mean/SD ve pooled OOF değerleri
    summary_df.to_excel(writer, sheet_name="Performance_Details", index=False)

    # Her outer fold train/test sonuçları
    fold_results_df.to_excel(writer, sheet_name="Outer_Fold_Details", index=False)

    # Her satırın held-out prediction'ı ve outer fold bilgisi
    oof_predictions_df.to_excel(writer, sheet_name="OOF_Predictions", index=False)

    # Outer fold'larda bulunan formüller: stability/sensitivity kontrolü için
    outer_formulas_df.to_excel(writer, sheet_name="Outer_Fold_Formulas", index=False)

    # Makalede kullanılacak final full-data formüller
    final_formulas_df.to_excel(writer, sheet_name="Final_Formulas", index=False)

    # Final PySR aramasındaki candidate equations
    if not final_pysr_equations_df.empty:
        final_pysr_equations_df.to_excel(writer, sheet_name="Final_PySR_Candidates", index=False)

    # Reviewer-oriented high-frequency diagnostic
    high_freq_df.to_excel(writer, sheet_name="High_Frequency_Diagnostics", index=False)

    # Public database / reproducibility: group IDs dahil
    df.to_excel(writer, sheet_name="Data_with_Groups", index=False)

    # Ayarlar
    settings_df.to_excel(writer, sheet_name="Settings", index=False)

# ======================================================
# 15) SAVE WORD REPORT
# ======================================================
word_path = os.path.join(output_dir, "Frekans_Formulleri_GroupCV_Raporu.docx")
doc.save(word_path)

print("\n" + "=" * 90)
print("TÜM İŞLEMLER TAMAMLANDI")
print("=" * 90)
print("Excel:", excel_path)
print("Word :", word_path)
print("\nMakale performans tablosu için Excel -> Paper_Ready_Performance sayfasını kullanın.")
print("Final denklemler için Excel -> Final_Formulas sayfasını kullanın.")
print("Parity plots pooled held-out geometry-group predictions üzerinden hazırlanmıştır.")
