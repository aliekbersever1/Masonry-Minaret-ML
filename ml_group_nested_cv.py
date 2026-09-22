# -*- coding: utf-8 -*-
"""
Geometry-grouped nested cross-validation for predicting the first two natural
frequencies (f1 and f2) of masonry minarets.

Public/reproducible version used for the manuscript.

Main features
-------------
- Predictors: H, D0, Dtop, t, B, and sqrt(E/rho).
- H is the only height-related predictor; Hc/H_silindir is not used.
- Geometry groups are defined by H, D0, Dtop, t, and B.
- Nested 5-fold GroupKFold: inner folds tune hyperparameters; outer folds
  provide held-out generalization estimates.
- Random Forest, XGBoost, LightGBM, and CatBoost.
- Target transformation: log1p / expm1 through TransformedTargetRegressor.
- Fold-wise mean +/- SD and pooled out-of-fold metrics.
- Parity plots with 1:1 line and +/-20% relative-error bounds.
- Held-out permutation importance and SHAP interpretation of final models.
- Final models are re-tuned on the complete dataset only after outer-CV
  evaluation and are saved for reproducibility/deployment.

Default repository layout
-------------------------
repo_root/
  data/minaret_frequency_database.xlsx
  results/ml/
  ml_group_nested_cv.py

The paths can be overridden with --data and --output.
"""

import os
import argparse
from pathlib import Path
import warnings
import traceback
import joblib
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import shap

from sklearn.model_selection import GroupKFold, RandomizedSearchCV
from sklearn.compose import TransformedTargetRegressor
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.ensemble import RandomForestRegressor
from sklearn.inspection import permutation_importance

from xgboost import XGBRegressor
from lightgbm import LGBMRegressor
from catboost import CatBoostRegressor

from docx import Document
from docx.shared import Inches

warnings.filterwarnings("ignore")
plt.ioff()

# ======================================================
# 1) PATHS / COMMAND-LINE OPTIONS
# ======================================================
def parse_args():
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(
        description="Geometry-grouped nested CV for masonry-minaret frequencies."
    )
    parser.add_argument(
        "--data",
        type=Path,
        default=script_dir / "data" / "minaret_frequency_database.xlsx",
        help="Input Excel database.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=script_dir / "results" / "ml",
        help="Directory for generated results.",
    )
    return parser.parse_args()


ARGS = parse_args()
data_path = ARGS.data.expanduser().resolve()
output_dir = ARGS.output.expanduser().resolve()

if not data_path.exists():
    raise FileNotFoundError(
        f"Input database not found: {data_path}\n"
        "Place the database at data/minaret_frequency_database.xlsx "
        "or provide a path with --data."
    )

output_dir.mkdir(parents=True, exist_ok=True)

# ======================================================
# 2) GENERAL SETTINGS
# ======================================================
RANDOM_STATE = 42
N_OUTER_SPLITS = 5
N_INNER_SPLITS = 5
RANDOM_SEARCH_ITER = 25
N_JOBS = -1
PERMUTATION_REPEATS = 20

USE_SHAP_SAMPLING = True
SHAP_SAMPLE_SIZE = 1200
SHAP_BAR_COLOR = "#ff005c"

# Aynı H, D0, Dtop, t ve B geometrisine ait E/rho varyasyonlarının train ve test'e dağılmasını engeller.
GEOMETRY_COLS = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
]

TARGET_COLS = ["f1 (Hz)", "f2 (Hz)"]

# ======================================================
# 3) HELPER FUNCTIONS
# ======================================================
def clean_numeric_series(series):
    return pd.to_numeric(
        series.astype(str)
        .str.replace(",", ".", regex=False)
        .str.replace("[", "", regex=False)
        .str.replace("]", "", regex=False)
        .str.strip(),
        errors="coerce",
    )


def clean_name(text):
    return (
        str(text)
        .replace(" ", "_")
        .replace("(", "")
        .replace(")", "")
        .replace("/", "_")
    )


def make_geometry_groups(dataframe, geometry_cols, decimals=8):
    """
    Stable G1...Gn geometry identifiers oluşturur.
    Grup tanımı yalnızca geometrik parametrelere dayanır.

    G etiketleri, yuvarlanmış benzersiz geometry tuple'larının sıralanmasıyla
    atanır; böylece dataframe satır sırası değişse bile grup adları kararlı kalır.
    """
    missing = [c for c in geometry_cols if c not in dataframe.columns]
    if missing:
        raise ValueError(
            "Geometry group oluşturmak için gerekli sütunlar eksik: "
            + ", ".join(missing)
        )

    geom = dataframe[geometry_cols].apply(pd.to_numeric, errors="coerce").round(decimals)
    if geom.isna().any().any():
        raise ValueError("Geometry columns içinde NaN var; group oluşturulamıyor.")

    tuples = [tuple(row) for row in geom.to_numpy()]
    unique_sorted = sorted(set(tuples))
    tuple_to_group = {key: f"G{i + 1}" for i, key in enumerate(unique_sorted)}

    groups = pd.Series(
        [tuple_to_group[key] for key in tuples],
        index=dataframe.index,
        name="Geometry_Group",
    )
    keys = pd.Series(
        ["_".join(f"{v:g}" for v in key) for key in tuples],
        index=dataframe.index,
        name="Geometry_Key",
    )

    return groups, keys


def smape_percent(y_true, y_pred):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    denom = np.abs(y_true) + np.abs(y_pred)
    valid = denom > 0
    if not np.any(valid):
        return 0.0
    return float(np.mean(2.0 * np.abs(y_pred[valid] - y_true[valid]) / denom[valid]) * 100.0)


def within_relative_band_percent(y_true, y_pred, band=0.20):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    valid = np.abs(y_true) > 1e-12
    if not np.any(valid):
        return np.nan
    rel_err = np.abs(y_pred[valid] - y_true[valid]) / np.abs(y_true[valid])
    return float(np.mean(rel_err <= band) * 100.0)


def calculate_metrics(y_true, y_pred):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)

    mae = mean_absolute_error(y_true, y_pred)
    rmse = np.sqrt(mean_squared_error(y_true, y_pred))
    r2 = r2_score(y_true, y_pred)
    smape = smape_percent(y_true, y_pred)
    within20 = within_relative_band_percent(y_true, y_pred, band=0.20)

    if np.std(y_true) == 0 or np.std(y_pred) == 0:
        r = np.nan
    else:
        r = np.corrcoef(y_true, y_pred)[0, 1]

    return {
        "R2": float(r2),
        "MAE": float(mae),
        "RMSE": float(rmse),
        "SMAPE_pct": float(smape),
        "CC": float(r) if pd.notna(r) else np.nan,
        "Within_20pct_pct": float(within20) if pd.notna(within20) else np.nan,
    }


def mean_sd(series):
    s = pd.Series(series, dtype=float)
    return float(s.mean()), float(s.std(ddof=1))


def get_valid_group_kfold(groups, requested_splits):
    n_groups = pd.Series(groups).nunique()
    n_splits = min(requested_splits, n_groups)
    if n_splits < 2:
        raise ValueError(f"GroupKFold için en az 2 group gerekli; bulunan group sayısı: {n_groups}")
    return GroupKFold(n_splits=n_splits)


# ======================================================
# 4) LOAD AND CLEAN DATA
# ======================================================
df = pd.read_excel(data_path)
df.columns = [str(c).strip() for c in df.columns]

print("İlk veri boyutu:", df.shape)
print("Sütunlar:", list(df.columns))

raw_needed_cols = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
    "E (MPa)",
    "Ro (kg/m3)",
    "f1 (Hz)",
    "f2 (Hz)",
]

missing_raw = [c for c in raw_needed_cols if c not in df.columns]
if missing_raw:
    raise ValueError("Eksik gerekli sütunlar: " + ", ".join(missing_raw))

for col in raw_needed_cols:
    df[col] = clean_numeric_series(df[col])

# E/rho fiziksel eğilimini kompakt predictor olarak koruyoruz.
df["sqrt_E_div_rho"] = np.sqrt(df["E (MPa)"] / df["Ro (kg/m3)"])

feature_cols = [
    "H_m",
    "D0_m",
    "Dtop_m",
    "t_m",
    "B_kaide_m",
    "sqrt_E_div_rho",
]

feature_label_map = {
    "H_m": r"$H$",
    "sqrt_E_div_rho": r"$\sqrt{E/\rho}$",
    "D0_m": r"$D_0$",
    "t_m": r"$t$",
    "B_kaide_m": r"$B$",
    "Dtop_m": r"$D_{top}$",
}

numeric_cols = list(dict.fromkeys(raw_needed_cols + ["sqrt_E_div_rho"]))
for col in numeric_cols:
    df[col] = clean_numeric_series(df[col])

# ML ve grouping için gerekli satırlar tam olmalı.
df = df.dropna(subset=list(dict.fromkeys(feature_cols + TARGET_COLS + GEOMETRY_COLS))).reset_index(drop=True)

# Stable geometry groups.
groups_full, geometry_keys = make_geometry_groups(df, GEOMETRY_COLS)
df["Geometry_Group"] = groups_full
df["Geometry_Key"] = geometry_keys

n_groups = groups_full.nunique()
if n_groups < N_OUTER_SPLITS:
    raise ValueError(
        f"Outer GroupKFold için {N_OUTER_SPLITS} fold istendi ancak yalnızca {n_groups} unique geometry group var."
    )

print("Temizlenmiş veri boyutu:", df.shape)
print("Unique geometry group sayısı:", n_groups)
print("Group definition:", ", ".join(GEOMETRY_COLS))

# ======================================================
# 5) MODELS AND HYPERPARAMETER SEARCH SPACES
# ======================================================
def build_models():
    models = {}

    rf_model = RandomForestRegressor(
        random_state=RANDOM_STATE,
        n_jobs=N_JOBS,
    )
    rf_ttr = TransformedTargetRegressor(
        regressor=rf_model,
        func=np.log1p,
        inverse_func=np.expm1,
    )
    rf_params = {
        "regressor__n_estimators": [100, 200, 300, 500, 800],
        "regressor__max_depth": [None, 4, 6, 8, 12, 16, 25],
        "regressor__min_samples_split": [2, 4, 8, 10],
        "regressor__min_samples_leaf": [1, 2, 3, 4],
        "regressor__max_features": ["sqrt", 0.6, 0.8, 1.0],
    }
    models["RandomForest"] = (rf_ttr, rf_params)

    xgb_model = XGBRegressor(
        objective="reg:squarederror",
        random_state=RANDOM_STATE,
        tree_method="hist",
        verbosity=0,
        n_jobs=N_JOBS,
    )
    xgb_ttr = TransformedTargetRegressor(
        regressor=xgb_model,
        func=np.log1p,
        inverse_func=np.expm1,
    )
    xgb_params = {
        "regressor__n_estimators": [100, 200, 300, 500, 800],
        "regressor__max_depth": [2, 3, 4, 5, 6, 8],
        "regressor__learning_rate": [0.01, 0.03, 0.05, 0.10, 0.15],
        "regressor__subsample": [0.6, 0.7, 0.8, 0.9, 1.0],
        "regressor__colsample_bytree": [0.6, 0.7, 0.8, 0.9, 1.0],
        "regressor__min_child_weight": [1, 2, 3, 5],
        "regressor__reg_alpha": [0.0, 0.001, 0.01, 0.1],
        "regressor__reg_lambda": [0.5, 1.0, 2.0, 5.0],
    }
    models["XGBoost"] = (xgb_ttr, xgb_params)

    lgbm_model = LGBMRegressor(
        random_state=RANDOM_STATE,
        verbose=-1,
        n_jobs=N_JOBS,
    )
    lgbm_ttr = TransformedTargetRegressor(
        regressor=lgbm_model,
        func=np.log1p,
        inverse_func=np.expm1,
    )
    lgbm_params = {
        "regressor__n_estimators": [100, 200, 300, 500, 800],
        "regressor__max_depth": [-1, 3, 4, 5, 6, 8, 10],
        "regressor__learning_rate": [0.01, 0.03, 0.05, 0.10, 0.15],
        "regressor__num_leaves": [7, 15, 31, 50, 70],
        "regressor__min_child_samples": [5, 10, 15, 20, 25],
        "regressor__subsample": [0.6, 0.7, 0.8, 0.9, 1.0],
        "regressor__colsample_bytree": [0.6, 0.8, 1.0],
        "regressor__reg_alpha": [0.0, 0.001, 0.01, 0.1],
        "regressor__reg_lambda": [0.5, 1.0, 2.0, 5.0],
    }
    models["LightGBM"] = (lgbm_ttr, lgbm_params)

    cat_model = CatBoostRegressor(
        random_seed=RANDOM_STATE,
        verbose=0,
        allow_writing_files=False,
    )
    cat_ttr = TransformedTargetRegressor(
        regressor=cat_model,
        func=np.log1p,
        inverse_func=np.expm1,
    )
    cat_params = {
        "regressor__iterations": [100, 200, 300, 500, 800],
        "regressor__depth": [3, 4, 5, 6, 7, 8],
        "regressor__learning_rate": [0.01, 0.03, 0.05, 0.10, 0.15],
        "regressor__l2_leaf_reg": [1, 3, 5, 7, 9, 12],
        "regressor__subsample": [0.6, 0.7, 0.8, 0.9, 1.0],
    }
    models["CatBoost"] = (cat_ttr, cat_params)

    return models


# ======================================================
# 6) GEOMETRY-GROUPED NESTED CROSS-VALIDATION
# ======================================================
def nested_group_cv_predictions(
    model,
    param_dist,
    X,
    y,
    groups,
    model_name,
    target_name,
):
    """
    Her outer fold için:
    1) GroupKFold ile geometry groups bazında train/test ayrımı.
    2) Inner GroupKFold sadece outer-training geometry groups üzerinde tuning.
    3) Outer-test geometry groups üzerinde tarafsız değerlendirme.
    4) Held-out outer-test üzerinde permutation importance.
    """
    outer_cv = GroupKFold(n_splits=N_OUTER_SPLITS)
    outer_splits = list(outer_cv.split(X, y, groups))

    oof_pred = np.full(len(y), np.nan, dtype=float)
    outer_fold_id = np.full(len(y), np.nan, dtype=float)
    fold_rows = []
    permutation_rows = []

    total_combinations = int(np.prod([len(v) for v in param_dist.values()]))
    n_iter = min(RANDOM_SEARCH_ITER, total_combinations)

    for fold_no, (train_idx, test_idx) in enumerate(outer_splits, start=1):
        X_train = X.iloc[train_idx].copy()
        X_test = X.iloc[test_idx].copy()
        y_train = y.iloc[train_idx].copy()
        y_test = y.iloc[test_idx].copy()
        groups_train = groups.iloc[train_idx].copy()
        groups_test = groups.iloc[test_idx].copy()

        train_group_set = set(groups_train.astype(str))
        test_group_set = set(groups_test.astype(str))
        overlap = train_group_set.intersection(test_group_set)
        if overlap:
            raise RuntimeError(
                f"Geometry-group leakage detected in outer fold {fold_no}: {sorted(overlap)}"
            )

        inner_cv = get_valid_group_kfold(groups_train, N_INNER_SPLITS)

        search = RandomizedSearchCV(
            estimator=model,
            param_distributions=param_dist,
            n_iter=n_iter,
            cv=inner_cv,
            scoring="r2",
            n_jobs=N_JOBS,
            pre_dispatch="2*n_jobs",
            random_state=RANDOM_STATE + fold_no,
            verbose=0,
            refit=True,
        )

        # groups=groups_train => inner CV de geometry-group based.
        search.fit(X_train, y_train, groups=groups_train)

        best_model = search.best_estimator_
        train_pred = best_model.predict(X_train)
        test_pred = best_model.predict(X_test)

        oof_pred[test_idx] = test_pred
        outer_fold_id[test_idx] = fold_no

        train_metrics = calculate_metrics(y_train, train_pred)
        test_metrics = calculate_metrics(y_test, test_pred)

        fold_rows.append({
            "Target": target_name,
            "Model": model_name,
            "Outer_Fold": fold_no,
            "Train_n": len(train_idx),
            "Test_n": len(test_idx),
            "Train_Groups": groups_train.nunique(),
            "Test_Groups": groups_test.nunique(),
            "Group_Overlap_n": len(overlap),
            "Inner_Best_R2": search.best_score_,
            "Train_R2": train_metrics["R2"],
            "Train_MAE_Hz": train_metrics["MAE"],
            "Train_RMSE_Hz": train_metrics["RMSE"],
            "Train_SMAPE_pct": train_metrics["SMAPE_pct"],
            "Test_R2": test_metrics["R2"],
            "Test_MAE_Hz": test_metrics["MAE"],
            "Test_RMSE_Hz": test_metrics["RMSE"],
            "Test_SMAPE_pct": test_metrics["SMAPE_pct"],
            "Test_Within_20pct_pct": test_metrics["Within_20pct_pct"],
            "Best_Params": str(search.best_params_),
        })

        # Reviewer-oriented interpretability: test-fold permutation importance.
        try:
            perm = permutation_importance(
                best_model,
                X_test,
                y_test,
                n_repeats=PERMUTATION_REPEATS,
                random_state=RANDOM_STATE + fold_no,
                scoring="r2",
                n_jobs=N_JOBS,
            )
            for feature, imp_mean, imp_sd in zip(
                X.columns,
                perm.importances_mean,
                perm.importances_std,
            ):
                permutation_rows.append({
                    "Target": target_name,
                    "Model": model_name,
                    "Outer_Fold": fold_no,
                    "Feature": feature,
                    "Permutation_Importance_Mean": imp_mean,
                    "Permutation_Importance_SD_Repeats": imp_sd,
                })
        except Exception as exc:
            print(
                f"Permutation importance failed: {target_name} | {model_name} | fold {fold_no} | {exc}"
            )

        print(
            f"{target_name} | {model_name} | Outer fold {fold_no}/{N_OUTER_SPLITS} | "
            f"train groups={groups_train.nunique()} | test groups={groups_test.nunique()} | "
            f"R²={test_metrics['R2']:.4f} | SMAPE={test_metrics['SMAPE_pct']:.2f}%"
        )

    if np.isnan(oof_pred).any() or np.isnan(outer_fold_id).any():
        raise RuntimeError(
            f"{target_name} - {model_name}: bazı OOF predictions/fold assignments eksik."
        )

    return (
        oof_pred,
        outer_fold_id.astype(int),
        pd.DataFrame(fold_rows),
        pd.DataFrame(permutation_rows),
    )


# ======================================================
# 7) NİHAİ MODELİ GROUP-BASED CV İLE TÜM VERİ ÜZERİNDE OPTİMİZE ET
# ======================================================
def fit_final_group_model(model, param_dist, X, y, groups):
    """
    Bu model performans raporlamak için kullanılmaz.
    Deployment ve final SHAP analizi içindir.
    Hyperparameter selection yine GroupKFold ile yapılır.
    """
    total_combinations = int(np.prod([len(v) for v in param_dist.values()]))
    n_iter = min(RANDOM_SEARCH_ITER, total_combinations)
    final_cv = get_valid_group_kfold(groups, N_INNER_SPLITS)

    final_search = RandomizedSearchCV(
        estimator=model,
        param_distributions=param_dist,
        n_iter=n_iter,
        cv=final_cv,
        scoring="r2",
        n_jobs=N_JOBS,
        pre_dispatch="2*n_jobs",
        random_state=RANDOM_STATE,
        verbose=0,
        refit=True,
    )

    final_search.fit(X, y, groups=groups)
    return (
        final_search.best_estimator_,
        final_search.best_params_,
        final_search.best_score_,
    )


# ======================================================
# 8) GRAFİK FONKSİYONLARI
# ======================================================
def save_parity_plot(
    y_true,
    y_pred,
    fold_ids,
    model_name,
    target_name,
    metrics,
    save_path,
):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    fold_ids = np.asarray(fold_ids, dtype=int)

    fig, ax = plt.subplots(figsize=(7.0, 7.0))

    # Outer folds ayrı renkte gösterilir; böylece grouped split görsel olarak da izlenebilir.
    unique_folds = sorted(np.unique(fold_ids))
    cmap = plt.get_cmap("tab10")
    for i, fold in enumerate(unique_folds):
        mask = fold_ids == fold
        ax.scatter(
            y_true[mask],
            y_pred[mask],
            alpha=0.72,
            s=34,
            label=f"Outer fold {fold}",
            color=cmap(i % 10),
        )

    min_val = min(np.min(y_true), np.min(y_pred))
    max_val = max(np.max(y_true), np.max(y_pred))
    lower = max(0.0, min_val * 0.95)
    upper = max_val * 1.05
    line_x = np.linspace(lower, upper, 250)

    # Ideal 1:1 line.
    ax.plot(line_x, line_x, linestyle="--", linewidth=1.6, label="1:1 line")

    # ±20% relative error bands (reviewer-requested diagnostic used in previous study).
    ax.plot(line_x, 1.20 * line_x, linestyle=":", linewidth=1.3, label="±20% bounds")
    ax.plot(line_x, 0.80 * line_x, linestyle=":", linewidth=1.3)

    text_box = (
        f"Pooled OOF R² = {metrics['R2']:.4f}\n"
        f"MAE = {metrics['MAE']:.4f} Hz\n"
        f"RMSE = {metrics['RMSE']:.4f} Hz\n"
        f"SMAPE = {metrics['SMAPE_pct']:.2f}%\n"
        f"Within ±20% = {metrics['Within_20pct_pct']:.1f}%"
    )

    ax.text(
        0.03,
        0.97,
        text_box,
        transform=ax.transAxes,
        ha="left",
        va="top",
        fontsize=13.5,
        bbox=dict(facecolor="white", alpha=0.88),
    )

    ax.set_xlabel(f"FE reference {target_name}", fontsize=15)
    ax.set_ylabel(f"Grouped outer-test prediction {target_name}", fontsize=15)
    ax.set_title(f"{target_name} – {model_name}", fontsize=16)
    ax.set_xlim(lower, upper)
    ax.set_ylim(lower, upper)
    ax.set_aspect("equal", adjustable="box")
    ax.tick_params(labelsize=13)
    ax.legend(fontsize=11, loc="lower right")

    fig.tight_layout()
    fig.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_permutation_importance_plot(perm_df, model_name, target_name, save_path):
    if perm_df.empty:
        return None

    agg = (
        perm_df.groupby("Feature")["Permutation_Importance_Mean"]
        .agg(["mean", "std"])
        .reset_index()
        .sort_values("mean", ascending=True)
    )

    labels = [feature_label_map.get(f, f) for f in agg["Feature"]]
    xerr = agg["std"].fillna(0.0).to_numpy()

    fig, ax = plt.subplots(figsize=(7.5, 5.5))
    ax.barh(labels, agg["mean"], xerr=xerr, alpha=0.85)
    ax.set_xlabel("Decrease in held-out outer-test R² after permutation", fontsize=14)
    ax.set_title(f"{target_name} – {model_name} permutation importance", fontsize=15)
    ax.tick_params(labelsize=13)
    fig.tight_layout()
    fig.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return save_path


def get_inner_tree_model(fitted_ttr_model):
    # TransformedTargetRegressor fit edildikten sonra gerçek estimator regressor_ içinde bulunur.
    if hasattr(fitted_ttr_model, "regressor_"):
        return fitted_ttr_model.regressor_
    return fitted_ttr_model


def get_shap_input(X_df):
    if (not USE_SHAP_SAMPLING) or len(X_df) <= SHAP_SAMPLE_SIZE:
        return X_df.copy()
    return X_df.sample(n=SHAP_SAMPLE_SIZE, random_state=RANDOM_STATE)


def save_custom_shap_bar(
    mean_abs_shap,
    feature_names,
    model_name,
    shap_bar_path,
):
    imp_df = pd.DataFrame({
        "feature": feature_names,
        "importance": mean_abs_shap,
    }).sort_values("importance", ascending=True)

    fig, ax = plt.subplots(figsize=(8, 5.5))
    bars = ax.barh(imp_df["feature"], imp_df["importance"], color=SHAP_BAR_COLOR)

    max_imp = max(float(imp_df["importance"].max()), 1e-12)
    for bar, val in zip(bars, imp_df["importance"]):
        ax.text(
            bar.get_width() + max_imp * 0.01,
            bar.get_y() + bar.get_height() / 2.0,
            f"{val:.3f}",
            va="center",
            ha="left",
            fontsize=12,
            color=SHAP_BAR_COLOR,
        )

    ax.set_xlabel(r"mean(|SHAP value|)", fontsize=14)
    ax.set_ylabel("Feature", fontsize=14)
    ax.set_title(f"{model_name} SHAP importance", fontsize=15)
    ax.tick_params(axis="both", labelsize=13)
    fig.tight_layout()
    fig.savefig(shap_bar_path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_shap_plots(
    final_model,
    X_df,
    model_name,
    target_name,
    out_dir,
):
    """
    SHAP final/deployment model davranışını açıklar.
    Bu analiz outer-test performance değerlendirmesi değildir.
    """
    shap_summary_path = os.path.join(
        out_dir,
        f"SHAP_summary_{target_name}_{model_name}.png",
    )
    shap_bar_path = os.path.join(
        out_dir,
        f"SHAP_bar_{target_name}_{model_name}.png",
    )

    X_shap = get_shap_input(X_df).copy()
    X_shap_labeled = X_shap.rename(columns=feature_label_map)

    tree_model = get_inner_tree_model(final_model)
    explainer = shap.TreeExplainer(tree_model)
    shap_values = explainer.shap_values(X_shap)

    if isinstance(shap_values, list):
        shap_values = shap_values[0]

    plt.figure(figsize=(8, 5.5))
    shap.summary_plot(
        shap_values,
        X_shap_labeled,
        show=False,
        plot_size=None,
    )
    plt.xlabel("SHAP value (final fitted model output scale)", fontsize=14)

    # SHAP'in oluşturduğu ana eksen ve colorbar dahil tüm yazı/sayıları
    # önceki boyutlardan yaklaşık +3 pt daha büyük çizdir.
    shap_fig = plt.gcf()
    for shap_ax in shap_fig.axes:
        shap_ax.tick_params(axis="both", labelsize=13)
        shap_ax.xaxis.label.set_size(14)
        shap_ax.yaxis.label.set_size(14)
        shap_ax.title.set_size(15)
        for txt in shap_ax.texts:
            if txt.get_fontsize() < 13:
                txt.set_fontsize(13)

    plt.tight_layout()
    plt.savefig(shap_summary_path, dpi=300, bbox_inches="tight")
    plt.close()

    mean_abs_shap = np.abs(shap_values).mean(axis=0)
    labeled_feature_names = [
        feature_label_map.get(col, col) for col in X_shap.columns
    ]
    save_custom_shap_bar(
        mean_abs_shap,
        labeled_feature_names,
        model_name,
        shap_bar_path,
    )

    return shap_summary_path, shap_bar_path


def create_combined_figure(
    image_paths,
    labels,
    output_path,
    main_title="",
):
    n = len(image_paths)
    if n == 0:
        return None

    cols = 2
    rows = int(np.ceil(n / cols))
    fig, axes = plt.subplots(rows, cols, figsize=(13, 6.2 * rows))
    axes_flat = np.atleast_1d(axes).flatten()

    for i, ax in enumerate(axes_flat):
        if i < n:
            img = plt.imread(image_paths[i])
            ax.imshow(img)
            ax.axis("off")
            ax.text(
                0.5,
                -0.045,
                labels[i],
                transform=ax.transAxes,
                ha="center",
                va="top",
                fontsize=14,
            )
        else:
            ax.axis("off")

    if main_title:
        fig.suptitle(main_title, fontsize=19)

    plt.tight_layout(rect=[0, 0.02, 1, 0.96])
    plt.savefig(output_path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return output_path


# ======================================================
# 9) HIGH-FREQUENCY TAIL DIAGNOSTIC
# ======================================================
def high_frequency_tail_metrics(y_true, y_pred, model_name, target_name, quantile=0.90):
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)

    threshold = float(np.quantile(y_true, quantile))
    mask = y_true >= threshold

    if not np.any(mask):
        return None

    yt = y_true[mask]
    yp = y_pred[mask]
    signed_error = yt - yp  # positive => underprediction

    return {
        "Target": target_name,
        "Model": model_name,
        "Quantile_Threshold": quantile,
        "Frequency_Threshold_Hz": threshold,
        "n_Cases": int(mask.sum()),
        "Mean_Signed_Error_ActualMinusPred_Hz": float(np.mean(signed_error)),
        "MAE_Hz": float(mean_absolute_error(yt, yp)),
        "RMSE_Hz": float(np.sqrt(mean_squared_error(yt, yp))),
        "SMAPE_pct": float(smape_percent(yt, yp)),
        "Within_20pct_pct": float(within_relative_band_percent(yt, yp, 0.20)),
    }


# ======================================================
# 10) WORD REPORT
# ======================================================
doc = Document()
doc.add_heading("Group-Based Nested Cross-Validation ile Frekans Tahmini", level=1)
doc.add_paragraph(
    "Model performance was evaluated using geometry-grouped nested five-fold cross-validation. "
    "All realizations sharing the same geometry were kept within the same fold. "
    "The outer loop was used only for unbiased performance evaluation, whereas hyperparameters "
    "were selected exclusively within the inner GroupKFold loop."
)
doc.add_paragraph(
    "Geometry group definition: " + ", ".join(GEOMETRY_COLS) + "."
)
doc.add_paragraph(
    f"The cleaned dataset contains {len(df)} realizations and {n_groups} unique geometry groups."
)
doc.add_paragraph(
    "Parity plots include the ideal 1:1 line and ±20% relative-error bounds. "
    "Permutation importance is evaluated only on held-out outer-test folds; SHAP is computed "
    "from the final full-data model and is used only for model interpretation."
)

# ======================================================
# 11) MAIN ANALYSIS LOOP
# ======================================================
all_summary_rows = []
all_fold_results = []
all_oof_predictions = []
all_final_params = []
all_permutation_rows = []
all_high_freq_rows = []

for target in TARGET_COLS:
    print("\n" + "=" * 90)
    print(f"{target} için GROUP-BASED nested CV başlıyor...")
    print("=" * 90)

    valid_mask = df[target].notna()
    target_df = df.loc[valid_mask].reset_index(drop=True)

    X = target_df[feature_cols].copy()
    y = target_df[target].copy()
    groups = target_df["Geometry_Group"].astype(str).copy()

    if groups.nunique() < N_OUTER_SPLITS:
        raise ValueError(
            f"{target}: yalnızca {groups.nunique()} unique group var; {N_OUTER_SPLITS} outer folds uygulanamaz."
        )

    models = build_models()
    target_results = []
    parity_paths = {}
    shap_summary_paths = {}
    shap_bar_paths = {}
    permutation_paths = {}

    doc.add_heading(f"{target} Sonuçları", level=2)

    for model_name, (model, param_dist) in models.items():
        print(f"\n{model_name} için group-based nested CV...")

        (
            y_pred_oof,
            fold_ids,
            fold_df,
            perm_df,
        ) = nested_group_cv_predictions(
            model=model,
            param_dist=param_dist,
            X=X,
            y=y,
            groups=groups,
            model_name=model_name,
            target_name=target,
        )

        pooled_metrics = calculate_metrics(y, y_pred_oof)

        # Fold mean ± SD: manuscript table için ana değerlendirme.
        tr_r2_m, tr_r2_sd = mean_sd(fold_df["Train_R2"])
        tr_mae_m, tr_mae_sd = mean_sd(fold_df["Train_MAE_Hz"])
        tr_rmse_m, tr_rmse_sd = mean_sd(fold_df["Train_RMSE_Hz"])

        te_r2_m, te_r2_sd = mean_sd(fold_df["Test_R2"])
        te_mae_m, te_mae_sd = mean_sd(fold_df["Test_MAE_Hz"])
        te_rmse_m, te_rmse_sd = mean_sd(fold_df["Test_RMSE_Hz"])
        te_smape_m, te_smape_sd = mean_sd(fold_df["Test_SMAPE_pct"])
        te_20_m, te_20_sd = mean_sd(fold_df["Test_Within_20pct_pct"])

        target_results.append({
            "Target": target,
            "Model": model_name,
            "Train_R2_Mean": tr_r2_m,
            "Train_R2_SD": tr_r2_sd,
            "Train_MAE_Mean_Hz": tr_mae_m,
            "Train_MAE_SD_Hz": tr_mae_sd,
            "Train_RMSE_Mean_Hz": tr_rmse_m,
            "Train_RMSE_SD_Hz": tr_rmse_sd,
            "Test_R2_Mean": te_r2_m,
            "Test_R2_SD": te_r2_sd,
            "Test_MAE_Mean_Hz": te_mae_m,
            "Test_MAE_SD_Hz": te_mae_sd,
            "Test_RMSE_Mean_Hz": te_rmse_m,
            "Test_RMSE_SD_Hz": te_rmse_sd,
            "Test_SMAPE_Mean_pct": te_smape_m,
            "Test_SMAPE_SD_pct": te_smape_sd,
            "Test_Within_20pct_Mean_pct": te_20_m,
            "Test_Within_20pct_SD_pct": te_20_sd,
            "Pooled_OOF_R2": pooled_metrics["R2"],
            "Pooled_OOF_MAE_Hz": pooled_metrics["MAE"],
            "Pooled_OOF_RMSE_Hz": pooled_metrics["RMSE"],
            "Pooled_OOF_SMAPE_pct": pooled_metrics["SMAPE_pct"],
            "Pooled_OOF_CC": pooled_metrics["CC"],
            "Pooled_OOF_Within_20pct_pct": pooled_metrics["Within_20pct_pct"],
        })

        all_fold_results.append(fold_df)
        if not perm_df.empty:
            all_permutation_rows.append(perm_df)

        pred_df = pd.DataFrame({
            "Target": target,
            "Model": model_name,
            "Row_Index": target_df.index,
            "Geometry_Group": groups.values,
            "Outer_Fold": fold_ids,
            "Actual_Hz": y.values,
            "Grouped_OuterTest_Predicted_Hz": y_pred_oof,
            "Residual_ActualMinusPred_Hz": y.values - y_pred_oof,
        })
        all_oof_predictions.append(pred_df)

        high_row = high_frequency_tail_metrics(
            y,
            y_pred_oof,
            model_name,
            target,
            quantile=0.90,
        )
        if high_row is not None:
            all_high_freq_rows.append(high_row)

        # --------------------------------------------------
        # A) Parity plot: pooled grouped outer-test predictions
        # --------------------------------------------------
        parity_path = os.path.join(
            output_dir,
            f"Parity_GroupNested_{clean_name(target)}_{model_name}.png",
        )
        save_parity_plot(
            y_true=y,
            y_pred=y_pred_oof,
            fold_ids=fold_ids,
            model_name=model_name,
            target_name=target,
            metrics=pooled_metrics,
            save_path=parity_path,
        )
        parity_paths[model_name] = parity_path

        # --------------------------------------------------
        # B) Outer-test permutation importance
        # --------------------------------------------------
        if not perm_df.empty:
            perm_path = os.path.join(
                output_dir,
                f"PermutationImportance_OuterTest_{clean_name(target)}_{model_name}.png",
            )
            save_permutation_importance_plot(
                perm_df,
                model_name,
                target,
                perm_path,
            )
            permutation_paths[model_name] = perm_path

        # --------------------------------------------------
        # C) Final full-data group-based model (deployment + SHAP)
        # --------------------------------------------------
        final_model, final_params, final_cv_score = fit_final_group_model(
            model=model,
            param_dist=param_dist,
            X=X,
            y=y,
            groups=groups,
        )

        final_model_path = os.path.join(
            output_dir,
            f"Final_GroupCV_{clean_name(target)}_{model_name}.pkl",
        )
        joblib.dump(final_model, final_model_path)

        all_final_params.append({
            "Target": target,
            "Model": model_name,
            "Final_GroupCV_Best_R2": final_cv_score,
            "Final_Best_Params": str(final_params),
            "Saved_Model": final_model_path,
        })

        # --------------------------------------------------
        # D) SHAP final fitted model üzerinde
        # --------------------------------------------------
        try:
            shap_summary_path, shap_bar_path = save_shap_plots(
                final_model=final_model,
                X_df=X,
                model_name=model_name,
                target_name=clean_name(target),
                out_dir=output_dir,
            )
            shap_summary_paths[model_name] = shap_summary_path
            shap_bar_paths[model_name] = shap_bar_path
        except Exception as exc:
            print(f"SHAP üretilemedi: {model_name} | {target} | {exc}")
            traceback.print_exc()

    # ------------------------------------------------------
    # Target summary
    # ------------------------------------------------------
    perf_df = pd.DataFrame(target_results).sort_values(
        by="Test_R2_Mean",
        ascending=False,
    ).reset_index(drop=True)
    all_summary_rows.append(perf_df)

    # Paper-ready compact table.
    paper_df = pd.DataFrame({
        "Model": perf_df["Model"],
        "Test R2 (mean ± SD)": [
            f"{m:.4f} ± {s:.4f}"
            for m, s in zip(perf_df["Test_R2_Mean"], perf_df["Test_R2_SD"])
        ],
        "Test MAE Hz (mean ± SD)": [
            f"{m:.4f} ± {s:.4f}"
            for m, s in zip(perf_df["Test_MAE_Mean_Hz"], perf_df["Test_MAE_SD_Hz"])
        ],
        "Test RMSE Hz (mean ± SD)": [
            f"{m:.4f} ± {s:.4f}"
            for m, s in zip(perf_df["Test_RMSE_Mean_Hz"], perf_df["Test_RMSE_SD_Hz"])
        ],
        "Test SMAPE % (mean ± SD)": [
            f"{m:.2f} ± {s:.2f}"
            for m, s in zip(perf_df["Test_SMAPE_Mean_pct"], perf_df["Test_SMAPE_SD_pct"])
        ],
        "Pooled OOF R2": [f"{v:.4f}" for v in perf_df["Pooled_OOF_R2"]],
        "Pooled within ±20% (%)": [
            f"{v:.1f}" for v in perf_df["Pooled_OOF_Within_20pct_pct"]
        ],
    })

    # Target-specific Excel.
    excel_target_path = os.path.join(
        output_dir,
        f"GroupNestedCV_sonuclar_{clean_name(target)}.xlsx",
    )

    target_fold_df = pd.concat(
        [x for x in all_fold_results if not x.empty and x["Target"].iloc[0] == target],
        ignore_index=True,
    )
    target_pred_df = pd.concat(
        [x for x in all_oof_predictions if not x.empty and x["Target"].iloc[0] == target],
        ignore_index=True,
    )
    target_perm_list = [
        x for x in all_permutation_rows
        if not x.empty and x["Target"].iloc[0] == target
    ]
    target_perm_df = pd.concat(target_perm_list, ignore_index=True) if target_perm_list else pd.DataFrame()

    with pd.ExcelWriter(excel_target_path, engine="openpyxl") as writer:
        paper_df.to_excel(writer, sheet_name="Paper_Ready", index=False)
        perf_df.to_excel(writer, sheet_name="Performance_Details", index=False)
        target_fold_df.to_excel(writer, sheet_name="Outer_Fold_Details", index=False)
        target_pred_df.to_excel(writer, sheet_name="OOF_Predictions", index=False)
        if not target_perm_df.empty:
            target_perm_df.to_excel(writer, sheet_name="Permutation_Importance", index=False)

    # ------------------------------------------------------
    # Combined figures – fixed model order for consistency
    # ------------------------------------------------------
    ordered_models = ["RandomForest", "XGBoost", "LightGBM", "CatBoost"]
    letters = ["a", "b", "c", "d"]

    parity_image_list = [parity_paths[m] for m in ordered_models if m in parity_paths]
    parity_label_list = [
        f"({letters[i]}) {m}" for i, m in enumerate(ordered_models) if m in parity_paths
    ]
    combined_parity_path = os.path.join(
        output_dir,
        f"Combined_GroupNested_Parity_{clean_name(target)}.png",
    )
    create_combined_figure(
        parity_image_list,
        parity_label_list,
        combined_parity_path,
        main_title=f"{target} – grouped outer-test parity plots",
    )

    perm_image_list = [permutation_paths[m] for m in ordered_models if m in permutation_paths]
    perm_label_list = [
        f"({letters[i]}) {m}" for i, m in enumerate(ordered_models) if m in permutation_paths
    ]
    combined_perm_path = None
    if perm_image_list:
        combined_perm_path = os.path.join(
            output_dir,
            f"Combined_OuterTest_PermutationImportance_{clean_name(target)}.png",
        )
        create_combined_figure(
            perm_image_list,
            perm_label_list,
            combined_perm_path,
            main_title=f"{target} – outer-test permutation importance",
        )

    shap_summary_image_list = [
        shap_summary_paths[m] for m in ordered_models if m in shap_summary_paths
    ]
    shap_summary_label_list = [
        f"({letters[i]}) {m}" for i, m in enumerate(ordered_models) if m in shap_summary_paths
    ]
    combined_shap_summary_path = None
    if shap_summary_image_list:
        combined_shap_summary_path = os.path.join(
            output_dir,
            f"Combined_SHAP_Summary_{clean_name(target)}.png",
        )
        create_combined_figure(
            shap_summary_image_list,
            shap_summary_label_list,
            combined_shap_summary_path,
            main_title=f"{target} – SHAP summary plots (final fitted models)",
        )

    shap_bar_image_list = [
        shap_bar_paths[m] for m in ordered_models if m in shap_bar_paths
    ]
    shap_bar_label_list = [
        f"({letters[i]}) {m}" for i, m in enumerate(ordered_models) if m in shap_bar_paths
    ]
    combined_shap_bar_path = None
    if shap_bar_image_list:
        combined_shap_bar_path = os.path.join(
            output_dir,
            f"Combined_SHAP_Bar_{clean_name(target)}.png",
        )
        create_combined_figure(
            shap_bar_image_list,
            shap_bar_label_list,
            combined_shap_bar_path,
            main_title=f"{target} – SHAP importance plots (final fitted models)",
        )

    # ------------------------------------------------------
    # Word report
    # ------------------------------------------------------
    doc.add_heading("Paper-Ready Outer-Test Performance", level=3)
    table = doc.add_table(rows=1, cols=len(paper_df.columns))
    hdr = table.rows[0].cells
    for j, col in enumerate(paper_df.columns):
        hdr[j].text = str(col)

    for _, row in paper_df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(paper_df.columns):
            cells[j].text = str(row[col])

    doc.add_paragraph(
        f"Best mean outer-test R² model: {perf_df.loc[0, 'Model']}"
    )
    if os.path.exists(combined_parity_path):
        doc.add_picture(combined_parity_path, width=Inches(6.5))

    if combined_perm_path and os.path.exists(combined_perm_path):
        doc.add_heading("Held-Out Outer-Test Permutation Importance", level=3)
        doc.add_picture(combined_perm_path, width=Inches(6.5))

    if combined_shap_summary_path and os.path.exists(combined_shap_summary_path):
        doc.add_heading("SHAP Summary Plots – Final Fitted Models", level=3)
        doc.add_picture(combined_shap_summary_path, width=Inches(6.5))

    if combined_shap_bar_path and os.path.exists(combined_shap_bar_path):
        doc.add_heading("SHAP Importance Plots – Final Fitted Models", level=3)
        doc.add_picture(combined_shap_bar_path, width=Inches(6.5))


# ======================================================
# 12) SAVE ALL OUTPUTS
# ======================================================
summary_df = pd.concat(all_summary_rows, ignore_index=True)
fold_results_df = pd.concat(all_fold_results, ignore_index=True)
oof_predictions_df = pd.concat(all_oof_predictions, ignore_index=True)
final_params_df = pd.DataFrame(all_final_params)
permutation_df = (
    pd.concat(all_permutation_rows, ignore_index=True)
    if all_permutation_rows
    else pd.DataFrame()
)
high_freq_df = pd.DataFrame(all_high_freq_rows)

# Reviewer-friendly group export.
group_export_cols = [
    "Geometry_Group",
    "Geometry_Key",
    *GEOMETRY_COLS,
    "E (MPa)",
    "Ro (kg/m3)",
    "sqrt_E_div_rho",
    "f1 (Hz)",
    "f2 (Hz)",
]
group_export = df[group_export_cols].copy()

all_results_path = os.path.join(
    output_dir,
    "Tum_GroupNestedCV_Sonuclari.xlsx",
)

with pd.ExcelWriter(all_results_path, engine="openpyxl") as writer:
    summary_df.to_excel(writer, sheet_name="Performance_Summary", index=False)
    fold_results_df.to_excel(writer, sheet_name="Outer_Fold_Details", index=False)
    oof_predictions_df.to_excel(writer, sheet_name="OOF_Predictions", index=False)
    final_params_df.to_excel(writer, sheet_name="Final_Model_Params", index=False)
    high_freq_df.to_excel(writer, sheet_name="HighFreq_Tail", index=False)
    group_export.to_excel(writer, sheet_name="Data_with_Groups", index=False)
    if not permutation_df.empty:
        permutation_df.to_excel(writer, sheet_name="Permutation_Importance", index=False)

word_path = os.path.join(
    output_dir,
    "GroupNestedCV_ML_Frekans_Raporu.docx",
)
doc.save(word_path)

print("\nTüm işlemler tamamlandı.")
print("Toplu Excel:", all_results_path)
print("Word raporu:", word_path)
print("Geometry groups:", n_groups)
print("Grup tanımı:", GEOMETRY_COLS)
